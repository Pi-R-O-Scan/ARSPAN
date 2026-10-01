"""Sauvegarde mensuelle : convertit docs/data.json en fichier Excel (.xlsx).

Un onglet par graphique, une ligne par horodatage, une colonne par courbe.
Sortie : sauvegardes/aven_de_noel_AAAA-MM-JJ.xlsx (date du jour, heure de Paris).
Variables facultatives : DATA_FILE (défaut docs/data.json), OUT_DIR (défaut sauvegardes).
"""
import json
import os
import re
import sys
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATA_FILE = os.environ.get("DATA_FILE", "docs/data.json")
OUT_DIR = os.environ.get("OUT_DIR", "sauvegardes")
PARIS = ZoneInfo("Europe/Paris")
FMT = "dd/mm/yyyy hh:mm:ss"


def sheet_name(title, used):
    name = re.sub(r"[\[\]:*?/\\]", " ", title or "Mesures").strip()[:31] or "Mesures"
    base, n = name, 2
    while name in used:
        name = f"{base[:28]} {n}"
        n += 1
    used.add(name)
    return name


def main():
    with open(DATA_FILE, encoding="utf-8") as fh:
        data = json.load(fh)
    charts = data.get("charts", {})
    if not charts:
        sys.exit("Aucun graphique dans data.json : pas de sauvegarde.")

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    head_fill = PatternFill("solid", fgColor="D9E8E4")
    total = 0
    for cid, chart in charts.items():
        ws = wb.create_sheet(sheet_name(chart.get("title"), used))
        series = chart.get("series", [])
        unit = chart.get("unit", "")
        head = ["Date et heure (Paris)", "Date et heure (UTC)"] + [
            f"{s.get('name') or chart.get('title')} ({unit})" for s in series]
        ws.append(head)
        by_ts = {}
        for k, s in enumerate(series):
            for t, v in s["points"]:
                by_ts.setdefault(t, {})[k] = v
        for t in sorted(by_ts):
            utc = datetime.fromtimestamp(t / 1000, timezone.utc)
            paris = utc.astimezone(PARIS)
            row = [paris.replace(tzinfo=None), utc.replace(tzinfo=None)]
            row += [by_ts[t].get(k) for k in range(len(series))]
            ws.append(row)
        total += len(by_ts)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).number_format = FMT
            ws.cell(r, 2).number_format = FMT
        ws.freeze_panes = "A2"
        for i, h in enumerate(head, 1):
            ws.column_dimensions[get_column_letter(i)].width = 21 if i <= 2 else max(16, min(40, len(h) + 2))

    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, "aven_de_noel_" + datetime.now(PARIS).strftime("%Y-%m-%d") + ".xlsx")
    wb.save(out)
    print(f"Écrit : {out} ({len(charts)} onglets, {total} horodatages)")


if __name__ == "__main__":
    main()

